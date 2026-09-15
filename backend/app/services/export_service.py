"""
导出服务：把真题解析结果导出为 Markdown / Word
"""
import io
from datetime import datetime
from typing import Dict


def _safe_filename(text: str, max_len: int = 30) -> str:
    """生成安全的文件名，去掉非法字符"""
    import re
    text = re.sub(r'[\\/:*?"<>|\n\r\t]', "_", text)
    return text[:max_len].strip() or "export"


# ============================================================
# Markdown 导出
# ============================================================
def export_to_markdown(data: Dict) -> tuple:
    """
    把真题解析结果导出为 Markdown。
    返回 (文件名, 文件内容的 bytes)
    """
    question = data.get("question_text", "面试真题")
    res = data.get("result", {})

    lines = []
    lines.append(f"# {question}\n")
    lines.append(f"> 导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}\n")

    if data.get("position"):
        lines.append(f"**岗位**：{data['position']}\n")
    if data.get("tech_stack"):
        lines.append(f"**技术栈**：{', '.join(data['tech_stack'])}\n")

    lines.append(f"**题目类型**：{res.get('question_type', '-')}\n")

    lines.append("\n## 🎯 核心考察点\n")
    for p in res.get("core_points", []):
        lines.append(f"- {p}")

    lines.append("\n## ✅ 参考答案\n")
    lines.append(res.get("reference_answer", ""))

    lines.append("\n## 📝 答题框架\n")
    for step in res.get("answer_framework", []):
        lines.append(f"- {step}")

    lines.append("\n## 🔄 可能的追问\n")
    for f in res.get("follow_up_questions", []):
        lines.append(f"- {f}")

    lines.append("\n## ⚠️ 常见误区\n")
    for m in res.get("common_mistakes", []):
        lines.append(f"- {m}")

    content = "\n".join(lines)
    filename = f"面试真题解析_{_safe_filename(question)}.md"
    return filename, content.encode("utf-8")


# ============================================================
# Word 导出
# ============================================================
def export_to_docx(data: Dict) -> tuple:
    """把真题解析结果导出为 Word"""
    from docx import Document
    from docx.shared import Pt, RGBColor

    question = data.get("question_text", "面试真题")
    res = data.get("result", {})

    doc = Document()

    # 标题
    title = doc.add_heading(question, level=1)

    # 元信息
    meta = doc.add_paragraph()
    meta.add_run(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}\n").italic = True
    if data.get("position"):
        meta.add_run(f"岗位：{data['position']}\n")
    if data.get("tech_stack"):
        meta.add_run(f"技术栈：{', '.join(data['tech_stack'])}\n")
    meta.add_run(f"题目类型：{res.get('question_type', '-')}")

    # 核心考察点
    doc.add_heading("🎯 核心考察点", level=2)
    for p in res.get("core_points", []):
        doc.add_paragraph(p, style="List Bullet")

    # 参考答案
    doc.add_heading("✅ 参考答案", level=2)
    doc.add_paragraph(res.get("reference_answer", ""))

    # 答题框架
    doc.add_heading("📝 答题框架", level=2)
    for step in res.get("answer_framework", []):
        doc.add_paragraph(step, style="List Bullet")

    # 追问
    doc.add_heading("🔄 可能的追问", level=2)
    for f in res.get("follow_up_questions", []):
        doc.add_paragraph(f, style="List Bullet")

    # 常见误区
    doc.add_heading("⚠️ 常见误区", level=2)
    for m in res.get("common_mistakes", []):
        doc.add_paragraph(m, style="List Bullet")

    # 保存到内存
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = f"面试真题解析_{_safe_filename(question)}.docx"
    return filename, buffer.getvalue()